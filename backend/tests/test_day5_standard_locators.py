"""Bounded non-PDF evidence must not borrow wording from sibling sections."""

from __future__ import annotations

import io
from typing import Any

import pytest
from docx import Document

from app.curriculum_intelligence.framework_structure import (
    FrameworkStructureError,
    FrameworkStructureService,
)
from app.curriculum_intelligence.service import CurriculumIntelligenceError
from app.curriculum_intelligence.standards_locators import (
    StandardsLocatorError,
    resolve_standard_locator,
)
from app.models.curriculum import Competency
from app.schemas.curriculum_intelligence import CompetencySpec, OfficialSourceManifestEntry
from app.source_intelligence.security import FetchedSource, SourceUrlFetcher
from tests.test_day4_official_evidence import _pdf
from tests.test_day4_official_evidence import evidence_service as evidence_service
from tests.test_day5_standard_evidence import context


def source_case(kind: str) -> tuple[bytes, str, str, str, str]:
    if kind == "pdf":
        return (
            _pdf(["Section A C-A First standard", "Section B C-B Second standard"]),
            "application/pdf",
            "standards.pdf",
            "PDF page 1",
            "PDF page 2",
        )
    if kind == "text":
        return (
            b"Section A\nC-A First standard\nSection B\nC-B Second standard",
            "text/plain",
            "standards.txt",
            "text lines 1-2",
            "text lines 3-4",
        )
    if kind == "html":
        return (
            b'<html><body><section id="a"><h2>Section A</h2><p>C-A First standard</p>'
            b'</section><section id="b"><h2>Section B</h2><p>C-B Second standard</p>'
            b"</section></body></html>",
            "text/html",
            "standards.html",
            "HTML #a",
            "HTML #b",
        )
    if kind == "json":
        return (
            b'{"sections":{"A":"C-A First standard","B":"C-B Second standard"}}',
            "application/json",
            "standards.json",
            "JSON pointer /sections/A",
            "JSON pointer /sections/B",
        )
    if kind == "csv":
        return (
            b"Section A,C-A,First standard\nSection B,C-B,Second standard\n",
            "text/csv",
            "standards.csv",
            "CSV row 1",
            "CSV row 2",
        )
    document = Document()
    document.add_paragraph("Section A C-A First standard")
    document.add_paragraph("Section B C-B Second standard")
    output = io.BytesIO()
    document.save(output)
    return (
        output.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "standards.docx",
        "DOCX paragraph 1",
        "DOCX paragraph 2",
    )


class FormatFetcher(SourceUrlFetcher):
    def __init__(self, content: bytes, mime: str, filename: str) -> None:
        self.content, self.mime, self.filename = content, mime, filename

    def fetch(self, url: str) -> FetchedSource:
        return FetchedSource(self.content, url, self.mime, self.filename)


def actual_context(service: Any, kind: str) -> tuple[Any, Any, Any, str, str]:
    framework, _, node_spec = context(service)
    content, mime, filename, wrong, right = source_case(kind)
    service.source_service.fetcher = FormatFetcher(content, mime, filename)
    entry = OfficialSourceManifestEntry(
        key=f"synthetic-bounded-standard-{kind}",
        source_type="official_authority",
        title="Synthetic bounded standards",
        url=f"https://synthetic.example.invalid/{filename}",
        authority="Synthetic test authority",
        document_type="academic_standards",
    )
    revision = service.ensure_manifest_sources(
        [entry], actor_id="synthetic-test", fetch_content=True, fallback_on_fetch_error=False
    )[entry.key]
    return framework, revision, node_spec, wrong, right


@pytest.mark.parametrize("kind", ["text", "html", "json", "csv", "docx"])
@pytest.mark.parametrize("entrypoint", ["generic", "incoming-node", "attached-record", "valid"])
def test_non_pdf_sections_bind_both_records(evidence_service, kind, entrypoint):
    framework, revision, node_spec, wrong, right = actual_context(evidence_service, kind)
    spec = CompetencySpec(
        code=f"bounded-{kind}",
        name="Second standard",
        official_text="Second standard",
        source_locator=wrong if entrypoint == "generic" else right,
        metadata_json={"official_code": "C-B"},
    )
    if entrypoint == "generic":
        with pytest.raises(CurriculumIntelligenceError, match="absent at exact locator"):
            evidence_service.upsert_competencies(
                framework=framework, revision=revision, specs=[spec]
            )
        return
    competency = evidence_service.upsert_competencies(
        framework=framework, revision=revision, specs=[spec]
    )[spec.code]
    assert evidence_service.session.get(Competency, competency.id) is competency
    if entrypoint == "attached-record":
        # Simulate a historical/corrupted record bypassing today's writer.
        competency.source_locator = wrong
        evidence_service.session.flush()
    incoming = node_spec.model_copy(
        update={
            "code": f"bounded-node-{kind}",
            "official_code": "C-B",
            "official_text": "Second standard",
            "competency_id": competency.id,
            "source_locator": wrong if entrypoint == "incoming-node" else right,
        }
    )
    structure = FrameworkStructureService(
        evidence_service.session, source_service=evidence_service.source_service
    )
    if entrypoint != "valid":
        with pytest.raises(FrameworkStructureError, match="absent at exact locator"):
            structure.upsert_nodes(framework=framework, revision=revision, specs=[incoming])
    else:
        node = structure.upsert_nodes(framework=framework, revision=revision, specs=[incoming])[
            incoming.code
        ]
        assert node.source_locator == right and node.source_revision_id == revision.id
        assert node.competency_id == competency.id


@pytest.mark.parametrize("kind", ["text", "html", "json", "csv", "docx"])
def test_freeform_section_label_never_searches_whole_document(kind):
    content, mime, _, _, _ = source_case(kind)
    with pytest.raises(StandardsLocatorError):
        resolve_standard_locator(content, mime, "Section A")


@pytest.mark.parametrize(
    "content",
    [
        b'<section id="a">C-A</section><section id="a">C-B</section>',
        b'<section id="a">C-A<section id="b">C-B</section>',
        b'<section id="a"><div>C-A</section><section>C-B</section>',
        b'<section id="a" id="b">C-A</section>',
    ],
)
def test_ambiguous_html_locator_fails_closed(content):
    with pytest.raises(StandardsLocatorError):
        resolve_standard_locator(content, "text/html", "HTML #a")


def test_html_does_not_include_attributes_or_hidden_script_wording():
    text = resolve_standard_locator(
        b'<section id="a" title="C-B"><script>C-B</script><p>C-A</p></section>'
        b'<section id="b">C-B</section>',
        "text/html",
        "HTML #a",
    )
    assert "C-A" in text and "C-B" not in text


@pytest.mark.parametrize("locator", ["JSON pointer /a", "JSON pointer /a~2"])
def test_json_duplicate_keys_and_invalid_escapes_fail_closed(locator):
    with pytest.raises(StandardsLocatorError):
        resolve_standard_locator(b'{"a":"C-A","a":"C-B"}', "application/json", locator)


def test_json_pointer_escape_selects_only_exact_unicode_value():
    assert (
        resolve_standard_locator(
            '{"a/b":{"~code":"తెలుగు C-A"},"other":"C-B"}'.encode(),
            "application/json",
            "JSON pointer /a~1b/~0code",
        )
        == "తెలుగు C-A"
    )


def test_csv_column_cannot_borrow_from_other_column():
    assert resolve_standard_locator(b"C-A,C-B\n", "text/csv", "CSV row 1 column 1") == "C-A"


@pytest.mark.parametrize("locator", ["text lines 0-1", "text lines 2-1", "text lines 1-9"])
def test_invalid_text_bounds_rejected(locator):
    with pytest.raises(StandardsLocatorError):
        resolve_standard_locator(b"C-A\nC-B", "text/plain", locator)


@pytest.mark.parametrize("kind", ["pdf", "text", "html", "json", "csv", "docx"])
@pytest.mark.parametrize("entrypoint", ["generic", "incoming-node", "attached-record"])
def test_standard_code_component_cannot_masquerade_as_complete_code(
    evidence_service, kind, entrypoint
):
    framework, revision, node_spec, _, locator = actual_context(evidence_service, kind)
    spec = CompetencySpec(
        code=f"code-boundary-{kind}",
        name="Second standard",
        official_text="Second standard",
        source_locator=locator,
        metadata_json={"official_code": "C" if entrypoint == "generic" else "C-B"},
    )
    if entrypoint == "generic":
        with pytest.raises(CurriculumIntelligenceError, match="complete identifier"):
            evidence_service.upsert_competencies(
                framework=framework, revision=revision, specs=[spec]
            )
        return
    competency = evidence_service.upsert_competencies(
        framework=framework, revision=revision, specs=[spec]
    )[spec.code]
    competency.metadata_json = {"official_code": "C"}
    evidence_service.session.flush()
    incoming = node_spec.model_copy(
        update={
            "code": f"code-boundary-node-{kind}",
            "official_code": "C" if entrypoint == "incoming-node" else None,
            "official_text": "Second standard",
            "competency_id": competency.id,
            "source_locator": locator,
        }
    )
    structure = FrameworkStructureService(
        evidence_service.session, source_service=evidence_service.source_service
    )
    with pytest.raises(FrameworkStructureError, match="complete identifier"):
        structure.upsert_nodes(framework=framework, revision=revision, specs=[incoming])


@pytest.mark.parametrize(
    "text,code,expected",
    [
        ("C-10", "C-1", False),
        ("C-1", "C-1", True),
        ("(C-1).", "C-1", True),
        ("C-1.0", "C-1", False),
        ("C-1-A", "C-1", False),
        ("C-1/2", "C-1", False),
        ("C-1:sub", "C-1", False),
        ("XC-1", "C-1", False),
        ("X.C-1", "C-1", False),
        ("X-C-1", "C-1", False),
        ("C-1_suffix", "C-1", False),
        ("C-1 and C-10", "C-1", True),
        ("క-10", "క-1", False),
        ("క-1", "క-1", True),
        ("कि", "क", False),
        ("C-1\u200dextra", "C-1", False),
        ("", "", False),
        ("text", " ", False),
    ],
)
def test_official_identifier_boundaries(text, code, expected):
    from app.curriculum_intelligence.standards_evidence import contains_complete_identifier

    assert contains_complete_identifier(text, code) is expected

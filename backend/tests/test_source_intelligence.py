from __future__ import annotations

import io
import json
from pathlib import Path

import httpx
import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.db.base import Base
from app.db.session import create_database_engine
from app.models import (
    CurriculumPack,
    CurriculumVersion,
    EducationFramework,
    ExamPack,
    ExamVersion,
    PolicyRule,
    Question,
    SourceAuditEvent,
    SourceRevision,
)
from app.models.enums import (
    PolicyScopeType,
    QuestionOrigin,
    QuestionType,
    SourceIngestionMethod,
    SourceRevisionStatus,
    SourceTrustTier,
    SourceType,
)
from app.schemas.source_intelligence import (
    ManualSourceRevisionInput,
    SourceProvenanceLinkInput,
    SourceRegistrationInput,
)
from app.source_intelligence.extractors import SourceExtractionError
from app.source_intelligence.security import (
    SourceFetchError,
    SourceUrlFetcher,
    UnsafeSourceUrl,
    validate_source_url,
)
from app.source_intelligence.service import (
    SourceGovernanceError,
    SourceIntelligenceService,
    SourceLifecycleError,
)
from app.source_intelligence.storage import LocalSourceStorage, SourceStorageError

PUBLIC_IP = "93.184.216.34"


def _public_resolver(_: str) -> list[str]:
    return [PUBLIC_IP]


def _session_and_service(
    tmp_path: Path,
    *,
    fetcher: SourceUrlFetcher | None = None,
) -> tuple[Session, SourceIntelligenceService]:
    database_url = f"sqlite:///{tmp_path / 'source-intelligence.db'}"
    engine = create_database_engine(database_url)
    Base.metadata.create_all(engine)
    session = Session(engine)
    service = SourceIntelligenceService(
        session,
        storage=LocalSourceStorage(tmp_path / "private-sources"),
        fetcher=fetcher,
    )
    return session, service


def _official_source(
    service: SourceIntelligenceService,
    *,
    source_type: SourceType = SourceType.OFFICIAL_SYLLABUS,
    title: str = "Synthetic official syllabus",
    url: str | None = None,
):
    return service.register_source(
        SourceRegistrationInput(
            source_type=source_type,
            title=title,
            url=url,
            authority="Synthetic Education Authority",
            country="IN",
            board_or_exam="SYNTH",
            academic_year="2026-27",
            copyright_classification="official-public-document",
            trust_tier=SourceTrustTier.OFFICIAL_PRIMARY,
        ),
        actor_id="founder",
        request_id="req-register",
    )


def _extract_diff_validate_approve_activate(
    service: SourceIntelligenceService,
    revision: SourceRevision,
) -> SourceRevision:
    if revision.status == SourceRevisionStatus.STAGED.value:
        service.extract_revision(revision.id, actor_id="engineer")
    service.create_diff(revision.id, actor_id="engineer")
    result = service.validate_revision(revision.id, actor_id="engineer")
    assert result.valid
    service.approve_revision(revision.id, actor_id="founder")
    return service.activate_revision(revision.id, actor_id="founder")


def _simple_pdf(text: str) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
    )
    stream = DecodedStreamObject()
    safe_text = text.replace("(", "[").replace(")", "]")
    stream.set_data(f"BT /F1 12 Tf 72 720 Td ({safe_text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def _simple_docx(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    document.add_table(rows=1, cols=2)
    document.tables[0].cell(0, 0).text = "A"
    document.tables[0].cell(0, 1).text = "B"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_source_governance_rejects_mismatched_trust(tmp_path: Path) -> None:
    session, service = _session_and_service(tmp_path)
    try:
        with pytest.raises(SourceGovernanceError):
            service.register_source(
                SourceRegistrationInput(
                    source_type=SourceType.COMPETITIVE_ANALYSIS,
                    title="Synthetic analysis",
                    trust_tier=SourceTrustTier.OFFICIAL_PRIMARY,
                )
            )

        analytical = service.register_source(
            SourceRegistrationInput(
                source_type=SourceType.COMPETITIVE_ANALYSIS,
                title="Synthetic analysis",
                trust_tier=SourceTrustTier.ANALYTICAL,
            )
        )
        assert analytical.trust_tier == SourceTrustTier.ANALYTICAL.value
    finally:
        session.close()
        create_database_engine.cache_clear()


@pytest.mark.parametrize(
    ("method", "content", "filename", "expected_fragment"),
    [
        (
            SourceIngestionMethod.PDF,
            _simple_pdf("Official PDF content"),
            "source.pdf",
            "Official PDF content",
        ),
        (
            SourceIngestionMethod.DOCX,
            _simple_docx("Official DOCX content"),
            "source.docx",
            "Official DOCX content",
        ),
        (
            SourceIngestionMethod.CSV,
            b"chapter,weight\nAlgebra,10\nGeometry,5\n",
            "source.csv",
            "Algebra",
        ),
        (
            SourceIngestionMethod.JSON,
            json.dumps({"subject": "Mathematics", "year": 2026}).encode(),
            "source.json",
            "Mathematics",
        ),
    ],
)
def test_upload_ingestion_extracts_supported_formats(
    tmp_path: Path,
    method: SourceIngestionMethod,
    content: bytes,
    filename: str,
    expected_fragment: str,
) -> None:
    session, service = _session_and_service(tmp_path)
    try:
        source = _official_source(service)
        revision = service.ingest_upload(
            source.id,
            method=method,
            content=content,
            filename=filename,
            actor_id="founder",
        )
        assert revision.status == SourceRevisionStatus.STAGED.value
        assert revision.storage_path
        assert (tmp_path / "private-sources" / revision.storage_path).exists()

        extracted = service.extract_revision(revision.id, actor_id="engineer")
        assert extracted.status == SourceRevisionStatus.EXTRACTED.value
        assert expected_fragment in (extracted.extracted_text or "")
    finally:
        session.close()
        create_database_engine.cache_clear()


@pytest.mark.parametrize(
    ("method", "filename", "error_code"),
    [
        (SourceIngestionMethod.PDF, "broken.pdf", "corrupt_pdf"),
        (SourceIngestionMethod.DOCX, "broken.docx", "corrupt_docx"),
    ],
)
def test_corrupt_documents_have_explicit_failed_state_and_can_retry(
    tmp_path: Path,
    method: SourceIngestionMethod,
    filename: str,
    error_code: str,
) -> None:
    session, service = _session_and_service(tmp_path)
    try:
        source = _official_source(service)
        revision = service.ingest_upload(
            source.id,
            method=method,
            content=b"not a valid document",
            filename=filename,
        )
        with pytest.raises(SourceExtractionError):
            service.extract_revision(revision.id)

        session.refresh(revision)
        assert revision.status == SourceRevisionStatus.FAILED.value
        assert revision.failure_reason
        assert error_code in revision.failure_reason

        retried = service.retry_revision(revision.id, actor_id="founder")
        assert retried.status == SourceRevisionStatus.STAGED.value
        assert retried.failure_reason is None
    finally:
        session.close()
        create_database_engine.cache_clear()


def test_manual_ingestion_is_metadata_only_and_duplicate_safe(tmp_path: Path) -> None:
    session, service = _session_and_service(tmp_path)
    try:
        source = _official_source(service)
        payload = ManualSourceRevisionInput(
            metadata={"notice": "Official metadata fallback", "year": "2026-27"},
            note="Authority site blocked automated retrieval",
        )
        first = service.ingest_manual(source.id, payload, actor_id="founder")
        second = service.ingest_manual(source.id, payload, actor_id="founder")

        assert first.id == second.id
        assert first.storage_path is None
        assert first.status == SourceRevisionStatus.EXTRACTED.value
        revision_count = session.scalar(
            select(func.count(SourceRevision.id)).where(SourceRevision.source_id == source.id)
        )
        assert revision_count == 1

        duplicate_events = session.scalar(
            select(func.count(SourceAuditEvent.id)).where(
                SourceAuditEvent.source_id == source.id,
                SourceAuditEvent.event_type == "duplicate_content_detected",
            )
        )
        assert duplicate_events == 1
    finally:
        session.close()
        create_database_engine.cache_clear()


def test_url_ingestion_with_mocked_http(tmp_path: Path) -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        assert str(request.url) == "https://official.example/syllabus.json"
        return httpx.Response(
            200,
            headers={"content-type": "application/json"},
            json={"authority": "Synthetic", "version": 1},
        )

    fetcher = SourceUrlFetcher(
        timeout_seconds=1,
        max_bytes=4096,
        max_redirects=2,
        resolver=_public_resolver,
        transport=httpx.MockTransport(handler),
    )
    session, service = _session_and_service(tmp_path, fetcher=fetcher)
    try:
        source = _official_source(
            service,
            url="https://official.example/syllabus.json",
        )
        revision = service.ingest_url(source.id, actor_id="founder")
        assert revision.ingestion_method == SourceIngestionMethod.URL.value
        assert revision.metadata_json["final_url"] == "https://official.example/syllabus.json"
        extracted = service.extract_revision(revision.id)
        assert '"version": 1' in (extracted.extracted_text or "")
    finally:
        session.close()
        create_database_engine.cache_clear()


def test_url_network_failure_is_audited_and_retryable(tmp_path: Path) -> None:
    def failing_handler(request: httpx.Request) -> httpx.Response:
        raise httpx.ConnectError("synthetic connection failure", request=request)

    failing_fetcher = SourceUrlFetcher(
        timeout_seconds=1,
        max_bytes=4096,
        max_redirects=1,
        resolver=_public_resolver,
        transport=httpx.MockTransport(failing_handler),
    )
    session, service = _session_and_service(tmp_path, fetcher=failing_fetcher)
    try:
        source = _official_source(
            service,
            url="https://official.example/source.json",
        )
        with pytest.raises(SourceFetchError):
            service.ingest_url(
                source.id,
                actor_id="founder",
                request_id="req-failed-fetch",
            )

        failed_event = session.scalar(
            select(SourceAuditEvent).where(
                SourceAuditEvent.source_id == source.id,
                SourceAuditEvent.event_type == "source_ingestion_failed",
            )
        )
        assert failed_event is not None
        assert failed_event.outcome == "failed"
        assert failed_event.payload_json["ingestion_method"] == SourceIngestionMethod.URL.value
        assert failed_event.payload_json["error_type"] == "SourceFetchError"
        assert "source URL retrieval failed" in failed_event.payload_json["reason"]

        service.fetcher = SourceUrlFetcher(
            timeout_seconds=1,
            max_bytes=4096,
            max_redirects=1,
            resolver=_public_resolver,
            transport=httpx.MockTransport(
                lambda request: httpx.Response(
                    200,
                    request=request,
                    headers={"content-type": "application/json"},
                    content=b'{"retry":"succeeded"}',
                )
            ),
        )
        revision = service.ingest_url(source.id, actor_id="founder")
        assert revision.status == SourceRevisionStatus.STAGED.value
    finally:
        session.close()
        create_database_engine.cache_clear()


def test_unsupported_url_content_has_explicit_failed_revision_state(tmp_path: Path) -> None:
    fetcher = SourceUrlFetcher(
        timeout_seconds=1,
        max_bytes=4096,
        max_redirects=1,
        resolver=_public_resolver,
        transport=httpx.MockTransport(
            lambda request: httpx.Response(
                200,
                request=request,
                headers={"content-type": "application/zip"},
                content=b"PK synthetic unsupported content",
            )
        ),
    )
    session, service = _session_and_service(tmp_path, fetcher=fetcher)
    try:
        source = _official_source(
            service,
            url="https://official.example/source.zip",
        )
        revision = service.ingest_url(source.id)
        with pytest.raises(SourceExtractionError, match="unsupported source content type"):
            service.extract_revision(revision.id)

        session.refresh(revision)
        assert revision.status == SourceRevisionStatus.FAILED.value
        assert revision.failure_reason
        assert "unsupported_source_format" in revision.failure_reason
    finally:
        session.close()
        create_database_engine.cache_clear()


def test_url_security_blocks_unsafe_targets_redirects_and_size() -> None:
    with pytest.raises(UnsafeSourceUrl):
        validate_source_url("file:///tmp/source.pdf")
    with pytest.raises(UnsafeSourceUrl):
        validate_source_url("http://127.0.0.1/private")
    with pytest.raises(UnsafeSourceUrl):
        validate_source_url(
            "https://official.example/source",
            resolver=lambda _: ["10.0.0.1"],
        )

    def redirect_handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(302, headers={"location": "http://127.0.0.1/secret"})

    redirect_fetcher = SourceUrlFetcher(
        timeout_seconds=1,
        max_bytes=1024,
        max_redirects=2,
        resolver=_public_resolver,
        transport=httpx.MockTransport(redirect_handler),
    )
    with pytest.raises(UnsafeSourceUrl):
        redirect_fetcher.fetch("https://official.example/source")

    size_fetcher = SourceUrlFetcher(
        timeout_seconds=1,
        max_bytes=10,
        max_redirects=0,
        resolver=_public_resolver,
        transport=httpx.MockTransport(
            lambda request: httpx.Response(
                200,
                request=request,
                headers={"content-type": "text/plain", "content-length": "20"},
                content=b"x" * 20,
            )
        ),
    )
    with pytest.raises(SourceFetchError, match="size limit"):
        size_fetcher.fetch("https://official.example/source")


def test_no_silent_update_and_pattern_drift_activation_flow(tmp_path: Path) -> None:
    session, service = _session_and_service(tmp_path)
    try:
        source = _official_source(service)

        first = service.ingest_upload(
            source.id,
            method=SourceIngestionMethod.JSON,
            content=b'{"syllabus":["algebra"]}',
            filename="syllabus.json",
        )
        first = _extract_diff_validate_approve_activate(service, first)
        first_checksum = first.checksum
        assert source.checksum == first_checksum

        second = service.ingest_upload(
            source.id,
            method=SourceIngestionMethod.JSON,
            content=b'{"syllabus":["algebra","geometry"]}',
            filename="syllabus.json",
        )
        service.extract_revision(second.id)
        diff = service.create_diff(second.id)

        assert diff.from_revision_id == first.id
        assert diff.checksum_changed
        assert diff.content_diff_json["pattern_drift_candidate"] is True

        session.refresh(first)
        session.refresh(source)
        assert first.status == SourceRevisionStatus.ACTIVE.value
        assert second.status == SourceRevisionStatus.EXTRACTED.value
        assert source.checksum == first_checksum

        with pytest.raises(SourceLifecycleError):
            service.activate_revision(second.id, actor_id="founder")

        validation = service.validate_revision(second.id)
        assert validation.valid
        service.approve_revision(second.id, actor_id="founder")
        activated = service.activate_revision(second.id, actor_id="founder")

        session.refresh(first)
        session.refresh(source)
        assert activated.status == SourceRevisionStatus.ACTIVE.value
        assert activated.active_slot == 1
        assert first.status == SourceRevisionStatus.SUPERSEDED.value
        assert first.active_slot is None
        assert source.checksum == second.checksum
    finally:
        session.close()
        create_database_engine.cache_clear()


def test_failed_activation_rolls_back_without_false_success_log(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    caplog: pytest.LogCaptureFixture,
) -> None:
    session, service = _session_and_service(tmp_path)
    try:
        source = _official_source(service)
        first = service.ingest_manual(
            source.id,
            ManualSourceRevisionInput(metadata={"version": 1}),
        )
        first = _extract_diff_validate_approve_activate(service, first)

        second = service.ingest_manual(
            source.id,
            ManualSourceRevisionInput(metadata={"version": 2}),
        )
        service.create_diff(second.id)
        assert service.validate_revision(second.id).valid
        service.approve_revision(second.id, actor_id="founder")

        caplog.clear()

        def fail_commit() -> None:
            raise RuntimeError("synthetic commit failure")

        monkeypatch.setattr(session, "commit", fail_commit)
        with pytest.raises(RuntimeError, match="synthetic commit failure"):
            service.activate_revision(second.id, actor_id="founder")

        session.expire_all()
        persisted_first = session.get(SourceRevision, first.id)
        persisted_second = session.get(SourceRevision, second.id)
        assert persisted_first is not None
        assert persisted_second is not None
        assert persisted_first.status == SourceRevisionStatus.ACTIVE.value
        assert persisted_first.active_slot == 1
        assert persisted_second.status == SourceRevisionStatus.APPROVED.value
        assert persisted_second.active_slot is None

        messages = "\n".join(record.getMessage() for record in caplog.records)
        assert "source_activated" not in messages
        assert "source_superseded" not in messages
    finally:
        session.rollback()
        session.close()
        create_database_engine.cache_clear()


def test_conflicting_candidate_activation_requires_rediff_and_reapproval(
    tmp_path: Path,
) -> None:
    session, service = _session_and_service(tmp_path)
    try:
        source = _official_source(service)
        first = service.ingest_manual(
            source.id,
            ManualSourceRevisionInput(metadata={"version": 1}),
        )
        first = _extract_diff_validate_approve_activate(service, first)

        second = service.ingest_manual(
            source.id,
            ManualSourceRevisionInput(metadata={"version": 2}),
        )
        service.create_diff(second.id)
        assert service.validate_revision(second.id).valid
        service.approve_revision(second.id, actor_id="founder")

        third = service.ingest_manual(
            source.id,
            ManualSourceRevisionInput(metadata={"version": 3}),
        )
        third_diff = service.create_diff(third.id)
        assert third_diff.from_revision_id == first.id
        assert service.validate_revision(third.id).valid
        service.approve_revision(third.id, actor_id="founder")

        service.activate_revision(second.id, actor_id="founder")

        with pytest.raises(SourceLifecycleError, match="candidate diff is stale"):
            service.activate_revision(third.id, actor_id="founder")

        session.refresh(second)
        session.refresh(third)
        assert second.status == SourceRevisionStatus.ACTIVE.value
        assert third.status == SourceRevisionStatus.APPROVED.value

        refreshed_diff = service.create_diff(third.id, actor_id="founder")
        session.refresh(third)
        assert refreshed_diff.from_revision_id == second.id
        assert third.status == SourceRevisionStatus.EXTRACTED.value
        assert third.approved_by is None
        assert third.approved_at is None

        assert service.validate_revision(third.id, actor_id="founder").valid
        service.approve_revision(third.id, actor_id="founder")
        service.activate_revision(third.id, actor_id="founder")

        session.refresh(second)
        session.refresh(third)
        assert second.status == SourceRevisionStatus.SUPERSEDED.value
        assert third.status == SourceRevisionStatus.ACTIVE.value

        refreshed_event = session.scalar(
            select(SourceAuditEvent).where(
                SourceAuditEvent.source_id == source.id,
                SourceAuditEvent.source_revision_id == third.id,
                SourceAuditEvent.event_type == "source_diff_refreshed",
            )
        )
        assert refreshed_event is not None
    finally:
        session.close()
        create_database_engine.cache_clear()


def test_single_active_revision_database_invariant(tmp_path: Path) -> None:
    session, service = _session_and_service(tmp_path)
    try:
        source = _official_source(service)
        first = service.ingest_manual(
            source.id,
            ManualSourceRevisionInput(metadata={"version": 1}),
        )
        service.create_diff(first.id)
        service.validate_revision(first.id)
        service.approve_revision(first.id, actor_id="founder")
        service.activate_revision(first.id, actor_id="founder")

        second = SourceRevision(
            source_id=source.id,
            revision_number=2,
            ingestion_method=SourceIngestionMethod.MANUAL.value,
            checksum="a" * 64,
            content_type="application/vnd.eduvijna.metadata+json",
            byte_size=1,
            retrieved_at=first.retrieved_at,
            extraction_status="succeeded",
            extracted_text="{}",
            status=SourceRevisionStatus.ACTIVE.value,
            active_slot=1,
            metadata_json={},
        )
        session.add(second)
        with pytest.raises(IntegrityError):
            session.commit()
    finally:
        session.rollback()
        session.close()
        create_database_engine.cache_clear()


def test_validation_requires_diff_and_storage_integrity(tmp_path: Path) -> None:
    session, service = _session_and_service(tmp_path)
    try:
        source = _official_source(service)
        revision = service.ingest_upload(
            source.id,
            method=SourceIngestionMethod.JSON,
            content=b'{"version":1}',
            filename="source.json",
        )
        service.extract_revision(revision.id)
        rejected = service.validate_revision(revision.id)
        assert not rejected.valid
        assert any("diff_created" in error for error in rejected.errors)

        retried = service.retry_revision(revision.id)
        service.extract_revision(retried.id)
        service.create_diff(retried.id)
        assert retried.storage_path
        stored_path = tmp_path / "private-sources" / retried.storage_path
        stored_path.write_bytes(b"tampered")

        rejected = service.validate_revision(retried.id)
        assert not rejected.valid
        assert any("storage_integrity" in error for error in rejected.errors)
    finally:
        session.close()
        create_database_engine.cache_clear()


def test_provenance_links_are_revision_aware_and_hard_policy_is_governed(
    tmp_path: Path,
) -> None:
    session, service = _session_and_service(tmp_path)
    try:
        source = _official_source(service)
        revision = service.ingest_manual(
            source.id,
            ManualSourceRevisionInput(metadata={"rule": "synthetic"}),
        )
        revision = _extract_diff_validate_approve_activate(service, revision)

        framework = EducationFramework(code="fw", name="Framework", country="IN")
        curriculum_pack = CurriculumPack(
            code="curr",
            name="Curriculum",
            country="IN",
            framework=framework,
        )
        curriculum_version = CurriculumVersion(
            curriculum_pack=curriculum_pack,
            version_code="2026",
            status="active",
        )
        exam_pack = ExamPack(code="exam", name="Exam", country="IN")
        exam_version = ExamVersion(
            exam_pack=exam_pack,
            version_code="2026",
            status="active",
        )
        question = Question(
            origin_type=QuestionOrigin.GENERATED.value,
            question_type=QuestionType.DESCRIPTIVE.value,
            stem_text="Synthetic question",
        )
        hard_policy = PolicyRule(
            scope_type=PolicyScopeType.EXAM.value,
            scope_code="exam",
            policy_key="paper.question_count",
            value_json={"value": 10},
            priority=100,
            metadata_json={"enforcement": "hard"},
        )
        session.add_all(
            [
                framework,
                curriculum_pack,
                curriculum_version,
                exam_pack,
                exam_version,
                question,
                hard_policy,
            ]
        )
        session.commit()

        for link in (
            SourceProvenanceLinkInput(
                revision_id=revision.id,
                curriculum_version_id=curriculum_version.id,
            ),
            SourceProvenanceLinkInput(
                revision_id=revision.id,
                exam_version_id=exam_version.id,
            ),
            SourceProvenanceLinkInput(
                revision_id=revision.id,
                question_id=question.id,
            ),
            SourceProvenanceLinkInput(
                revision_id=revision.id,
                policy_rule_id=hard_policy.id,
            ),
        ):
            service.link_provenance(link, actor_id="founder")

        session.refresh(revision)
        assert curriculum_version in revision.curriculum_versions
        assert exam_version in revision.exam_versions
        assert question in revision.questions
        assert hard_policy in revision.policy_rules

        analytical = service.register_source(
            SourceRegistrationInput(
                source_type=SourceType.COMPETITIVE_ANALYSIS,
                title="Synthetic competitor statistics",
                trust_tier=SourceTrustTier.ANALYTICAL,
            )
        )
        analytical_revision = service.ingest_manual(
            analytical.id,
            ManualSourceRevisionInput(metadata={"frequency": 0.3}),
        )
        analytical_revision = _extract_diff_validate_approve_activate(
            service,
            analytical_revision,
        )

        with pytest.raises(SourceGovernanceError):
            service.link_provenance(
                SourceProvenanceLinkInput(
                    revision_id=analytical_revision.id,
                    policy_rule_id=hard_policy.id,
                )
            )
    finally:
        session.close()
        create_database_engine.cache_clear()


def test_audit_events_cover_lifecycle_without_storing_document_content(tmp_path: Path) -> None:
    session, service = _session_and_service(tmp_path)
    try:
        source = _official_source(service)
        revision = service.ingest_upload(
            source.id,
            method=SourceIngestionMethod.JSON,
            content=b'{"secret_document_text":"do-not-log-full-content"}',
            filename="source.json",
            request_id="req-1",
        )
        service.extract_revision(revision.id, request_id="req-2")
        service.create_diff(revision.id, request_id="req-3")
        service.validate_revision(revision.id, request_id="req-4")
        service.approve_revision(revision.id, actor_id="founder", request_id="req-5")
        service.activate_revision(revision.id, actor_id="founder", request_id="req-6")

        events = list(
            session.scalars(
                select(SourceAuditEvent)
                .where(SourceAuditEvent.source_id == source.id)
                .order_by(SourceAuditEvent.created_at)
            )
        )
        event_types = {event.event_type for event in events}
        assert {
            "source_registered",
            "source_ingested",
            "source_extracted",
            "source_diff_created",
            "source_checksum_verified",
            "source_validated",
            "source_approved",
            "source_activated",
        }.issubset(event_types)
        serialized = json.dumps([event.payload_json for event in events])
        assert "do-not-log-full-content" not in serialized
    finally:
        session.close()
        create_database_engine.cache_clear()


def test_private_storage_sanitizes_and_blocks_path_escape(tmp_path: Path) -> None:
    storage = LocalSourceStorage(tmp_path / "private")
    relative = storage.write_revision(
        source_id="source-1",
        revision_number=1,
        checksum="b" * 64,
        filename="../../dangerous syllabus?.json",
        content=b"{}",
    )
    assert ".." not in relative
    assert (tmp_path / "private" / relative).read_bytes() == b"{}"

    with pytest.raises(SourceStorageError):
        storage.read("../outside.txt")

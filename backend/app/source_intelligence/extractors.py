from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader


class SourceExtractionError(RuntimeError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True, slots=True)
class ExtractionResult:
    text: str
    metadata: dict[str, Any]


def normalize_text(value: str) -> str:
    lines = [line.rstrip() for line in value.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    normalized: list[str] = []
    blank = False
    for line in lines:
        if line.strip():
            normalized.append(line)
            blank = False
        elif not blank and normalized:
            normalized.append("")
            blank = True
    return "\n".join(normalized).strip()


def _decode_utf8(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise SourceExtractionError(
            "unsupported_text_encoding",
            "source text must be UTF-8 encoded",
        ) from exc


class _VisibleHtmlParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._skip_depth = 0
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        if tag in {"script", "style", "noscript"}:
            self._skip_depth += 1
        elif tag in {"p", "div", "li", "br", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript"} and self._skip_depth:
            self._skip_depth -= 1
        elif tag in {"p", "div", "li", "tr"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip_depth == 0:
            self.parts.append(data)


def _extract_pdf(content: bytes) -> ExtractionResult:
    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception as exc:
        raise SourceExtractionError("corrupt_pdf", "PDF could not be opened") from exc

    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception as exc:
            raise SourceExtractionError("encrypted_pdf", "PDF is encrypted") from exc
        if not unlocked:
            raise SourceExtractionError("encrypted_pdf", "PDF is encrypted")

    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:
            raise SourceExtractionError(
                "pdf_text_extraction_failed",
                "PDF text extraction failed",
            ) from exc

    text = normalize_text("\n\n".join(pages))
    if not text:
        raise SourceExtractionError(
            "pdf_no_extractable_text",
            "PDF contains no extractable text; OCR is not part of Day 3",
        )
    return ExtractionResult(text=text, metadata={"page_count": len(reader.pages)})


def _extract_docx(content: bytes) -> ExtractionResult:
    try:
        document = Document(io.BytesIO(content))
    except Exception as exc:
        raise SourceExtractionError("corrupt_docx", "DOCX could not be opened") from exc

    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    table_count = 0
    for table in document.tables:
        table_count += 1
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))

    text = normalize_text("\n".join(parts))
    if not text:
        raise SourceExtractionError("docx_no_text", "DOCX contains no extractable text")
    return ExtractionResult(
        text=text,
        metadata={
            "paragraph_count": len(document.paragraphs),
            "table_count": table_count,
        },
    )


def _extract_csv(content: bytes) -> ExtractionResult:
    decoded = _decode_utf8(content)
    try:
        rows = list(csv.reader(io.StringIO(decoded), strict=True))
    except csv.Error as exc:
        raise SourceExtractionError("invalid_csv", "CSV could not be parsed") from exc

    if not rows:
        raise SourceExtractionError("empty_csv", "CSV contains no rows")

    output = io.StringIO()
    writer = csv.writer(output, lineterminator="\n")
    writer.writerows(rows)
    return ExtractionResult(
        text=output.getvalue().strip(),
        metadata={
            "row_count": len(rows),
            "max_columns": max((len(row) for row in rows), default=0),
        },
    )


def _extract_json(content: bytes) -> ExtractionResult:
    decoded = _decode_utf8(content)
    def reject_constant(value: str) -> None:
        raise ValueError(f"non-standard JSON constant: {value}")

    try:
        payload = json.loads(decoded, parse_constant=reject_constant)
    except (json.JSONDecodeError, ValueError) as exc:
        raise SourceExtractionError("invalid_json", "JSON could not be parsed") from exc
    normalized = json.dumps(payload, ensure_ascii=False, sort_keys=True, indent=2)
    return ExtractionResult(
        text=normalized,
        metadata={"json_root_type": type(payload).__name__},
    )


def _extract_html(content: bytes) -> ExtractionResult:
    decoded = _decode_utf8(content)
    parser = _VisibleHtmlParser()
    try:
        parser.feed(decoded)
    except Exception as exc:
        raise SourceExtractionError("invalid_html", "HTML could not be parsed") from exc
    text = normalize_text("".join(parser.parts))
    if not text:
        raise SourceExtractionError("html_no_text", "HTML contains no extractable text")
    return ExtractionResult(text=text, metadata={"format": "html"})


def _extract_text(content: bytes) -> ExtractionResult:
    text = normalize_text(_decode_utf8(content))
    if not text:
        raise SourceExtractionError("empty_text", "source contains no text")
    return ExtractionResult(text=text, metadata={"format": "text"})


_MIME_FORMATS: dict[str, str] = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/csv": "csv",
    "application/csv": "csv",
    "application/json": "json",
    "text/json": "json",
    "text/html": "html",
    "application/xhtml+xml": "html",
    "text/plain": "text",
    "text/markdown": "text",
}

_SUFFIX_FORMATS: dict[str, str] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".csv": "csv",
    ".json": "json",
    ".html": "html",
    ".htm": "html",
    ".txt": "text",
    ".md": "text",
}


def _resolve_source_format(content_type: str, filename: str | None) -> str:
    mime = content_type.split(";", 1)[0].strip().lower()
    suffix = Path(filename or "").suffix.lower()
    mime_format = _MIME_FORMATS.get(mime)
    suffix_format = _SUFFIX_FORMATS.get(suffix)

    if mime_format is not None and suffix_format is not None and mime_format != suffix_format:
        raise SourceExtractionError(
            "source_format_mismatch",
            "source MIME type and filename extension disagree",
        )

    source_format = mime_format or suffix_format
    if source_format is None:
        raise SourceExtractionError(
            "unsupported_source_format",
            f"unsupported source content type: {content_type}",
        )
    return source_format


def extract_content(
    *,
    content: bytes,
    content_type: str,
    filename: str | None,
) -> ExtractionResult:
    source_format = _resolve_source_format(content_type, filename)

    if source_format == "pdf":
        return _extract_pdf(content)
    if source_format == "docx":
        return _extract_docx(content)
    if source_format == "csv":
        return _extract_csv(content)
    if source_format == "json":
        return _extract_json(content)
    if source_format == "html":
        return _extract_html(content)
    return _extract_text(content)
